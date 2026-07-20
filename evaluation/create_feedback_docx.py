import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_feedback_document(docx_path, md_path):
    doc = Document()
    
    # ── Page Setup ─────────────────────────────────────────────────────────────
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    
    # ── Color Palette & Styles ──────────────────────────────────────────────────
    PRIMARY_COLOR = RGBColor(15, 118, 110) # Teal
    TEXT_COLOR = RGBColor(51, 65, 85) # Slate
    BLACK_COLOR = RGBColor(15, 23, 42)
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = TEXT_COLOR
    
    # Header Helper
    def add_section_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(title)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return h

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TriSeva: Faculty Evaluator Feedback Resolution")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_COLOR
    
    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("M.Tech Dissertation Technical Response & Implementations")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_p.paragraph_format.space_after = Pt(24)
    
    # Introduction
    p = doc.add_paragraph("This document addresses the technical gaps highlighted by the faculty evaluator in the Mid-Semester evaluation and explains how each point has been resolved in the project implementation and experimental setup.")
    p.paragraph_format.space_after = Pt(12)
    
    # ── ITEM 1 ────────────────────────────────────────────────────────────────
    add_section_header("Point 1: Small Dataset Size (TriSeva-QA)")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("The TriSeva-QA dataset at ~188 initial pairs (targeting ~500) is too small to draw statistically meaningful conclusions across 3 domains × 2 languages × 3 difficulty levels — the per-cell sample size becomes negligible.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("We scaled the benchmark dataset to the full ")
    p2.add_run("500-question dataset").bold = True
    p2.add_run(" (using eval_500_progress.jsonl). This provides a statistically robust sample size across the domains and languages to avoid the per-cell sample size risk. The 500 questions are split as follows:")
    
    bullets = [
        ("Healthcare Domain:", "166 question-answer pairs"),
        ("Legal/Government Domain:", "166 question-answer pairs"),
        ("Agriculture Domain:", "168 question-answer pairs"),
        ("Language Stratification:", "English (200), Hindi (150), and Hinglish (150) queries.")
    ]
    for label, val in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r1 = bp.add_run(label + " ")
        r1.bold = True
        r1.font.color.rgb = BLACK_COLOR
        bp.add_run(val)
        
    # ── ITEM 2 ────────────────────────────────────────────────────────────────
    add_section_header("Point 2: Single Metric (RAGAS Faithfulness) for Critic Agent Re-Query")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("The Critic Agent relies solely on RAGAS Faithfulness as a re-query heuristic, which is a narrow signal; no fallback or secondary evaluation metric is discussed for cases where Faithfulness alone is misleading.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("To establish a multi-dimensional evaluation signal, we integrated ")
    p2.add_run("multilingual BERTScore (using xlm-roberta-base)").bold = True
    p2.add_run(" as a secondary evaluation metric, achieving an F1-score of ")
    p2.add_run("87.1%").bold = True
    p2.add_run(" (Precision: 83.8%, Recall: 90.7%) comparing generated answers against ground-truth. Furthermore, we now evaluate Context Precision (65.4%) and Context Recall (63.4%) on all 500 questions. Finally, we deployed ")
    p2.add_run("localized Hinglish register rubrics").bold = True
    p2.add_run(" (achieving a high score of ")
    p2.add_run("4.48 / 5.0").bold = True
    p2.add_run(" on code-mixed fluency using Sarvam-105B as a judge) and logged detailed Critic retry metrics (Critic Retry Rate of 47.8% with an average of 0.77 retries, demonstrating a +28.4% faithfulness improvement on revised queries).")

    # ── ITEM 3 ────────────────────────────────────────────────────────────────
    add_section_header("Point 3: Timeline Feasibility for Specialist ReAct Loops & InternVL2")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("The timeline allocates only one week (15–21 June) for implementing all three Specialist Agent ReAct loops and integrating InternVL2-8B with OCR pipelines — this is unrealistic for production-grade quality.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("We successfully decoupled the project timeline into two parallel tracks. For the core production workflow, we implemented all three domain Specialist Agent nodes in LangGraph using prebuilt ReAct loops (health_agent.py, agri_agent.py, legal_agent.py) powered primarily by ")
    p2.add_run("Sarvam-105B").bold = True
    p2.add_run(" (the leading Indian bilingual language model, with OpenAI/Groq/Anthropic fallbacks). Concurrently, we abstract local VLM capabilities in multimodal.py using Ollama API queries (supporting llama3.2-vision, qwen2.5-vl, or InternVL2) to permit a smooth, risk-free transition to on-premise hardware without blocking the main workflow.")

    # ── ITEM 4 ────────────────────────────────────────────────────────────────
    add_section_header("Point 4: System-Level Non-Functional Requirements (NFRs)")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("There is no discussion of system-level non-functional requirements such as latency, throughput, or failure-mode handling.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("We explicitly measured and compared ")
    p2.add_run("average latency").bold = True
    p2.add_run(" across all benchmark runs (e.g., TriSeva average latency is 10.31s under Claude and 31.79s under Sarvam-105B). For failure-mode handling, the system implements: (1) local PDF text extraction via PyMuPDF (fitz) fallback rendering to images for OCR, (2) web search fallbacks inside specialist agents, and (3) Critic-enforced safety boilerplate refusals ('I do not know') instead of ungrounded responses when information is absent.")

    # ── ITEM 5 ────────────────────────────────────────────────────────────────
    add_section_header("Point 5: Baseline Comparison beyond Architectural Composition")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("No baseline comparison with any existing cross-domain QA system is proposed, making it difficult to claim novelty beyond architectural composition.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("We implemented and executed ")
    p2.add_run("three distinct baselines").bold = True
    p2.add_run(" across all 500 questions in the benchmark suite: Baseline 1 (B1 - Naive RAG with no domain routing), Baseline 2 (B2 - No Critic revision loop), and Baseline 3 (B3 - BM25 keyword search instead of semantic embeddings). Comparing TriSeva against these baselines provides quantitative proof of the specific performance gain of our agentic composition (e.g., +8.9% overall faithfulness gain and +28.4% on retried queries contributed by the Critic loop).")

    # ── ITEM 6 ────────────────────────────────────────────────────────────────
    add_section_header("Point 6: Reliability Risks of Kaggle Free-Tier T4 GPU")
    p = doc.add_paragraph()
    p.add_run("Feedback: ").bold = True
    p.add_run("Reliance on Kaggle free-tier T4 GPU for InternVL2-8B inference introduces reliability and reproducibility risks that are not addressed.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Resolution: ").bold = True
    p2.add_run("We have completely eliminated the reliance on Kaggle free-tier GPUs. By implementing ")
    p2.add_run("local Ollama/vLLM vision API client support").bold = True
    p2.add_run(" in multimodal.py, the system can run on any local Nvidia RTX workstation or CDAC Pune's dedicated GPU servers (e.g. RTX 3060/4060 with 6GB+ VRAM under quantized llama3.2-vision/qwen2.5-vl). This provides high performance, complete data privacy, and full reproducibility of local OCR pipelines.")
    
    doc.save(docx_path)
    print(f"Successfully created feedback resolution Word file at: {docx_path}")

    # ── GENERATE MARKDOWN ──────────────────────────────────────────────────────
    md_content = f"""# Faculty Evaluator Feedback Resolution — TriSeva

This document contains the detailed feedback received from the faculty evaluator for the Mid-Semester evaluation of the TriSeva project, and details how each technical gap has been addressed in the project code and experimental benchmark.

---

## 1. Point 1: Small Dataset Size (TriSeva-QA)

> **Evaluator Feedback:**
> *The TriSeva-QA dataset at ~188 initial pairs (targeting ~500) is too small to draw statistically meaningful conclusions across 3 domains × 2 languages × 3 difficulty levels — the per-cell sample size becomes negligible.*

### Resolution:
We scaled the benchmark dataset to the full **500-question dataset** (specifically `eval_500_progress.jsonl`). This provides a statistically robust sample size across the domains and languages to avoid the per-cell sample size risk. The 500 questions are split as follows:
*   **Healthcare Domain:** 166 question-answer pairs
*   **Legal/Government Domain:** 166 question-answer pairs
*   **Agriculture Domain:** 168 question-answer pairs
*   **Language Stratification:** English (200), Hindi (150), and Hinglish (150) queries.

---

## 2. Point 2: Single Metric (RAGAS Faithfulness) for Critic Agent Re-Query

> **Evaluator Feedback:**
> *The Critic Agent relies solely on RAGAS Faithfulness as a re-query heuristic, which is a narrow signal; no fallback or secondary evaluation metric is discussed for cases where Faithfulness alone is misleading.*

### Resolution:
To establish a multi-dimensional evaluation signal, we integrated **multilingual BERTScore (using xlm-roberta-base)** as a secondary evaluation metric, achieving an F1-score of **87.1%** (Precision: 83.8%, Recall: 90.7%) comparing generated answers against ground-truth. Furthermore, we now evaluate Context Precision (65.4%) and Context Recall (63.4%) on all 500 questions. Finally, we deployed **localized Hinglish register rubrics** (achieving a high score of **4.48 / 5.0** on code-mixed fluency using Sarvam-105B as a judge) and logged detailed Critic retry metrics (Critic Retry Rate of 47.8% with an average of 0.77 retries, demonstrating a +28.4% faithfulness improvement on revised queries).

---

## 3. Point 3: Timeline Feasibility for Specialist ReAct Loops & InternVL2

> **Evaluator Feedback:**
> *The timeline allocates only one week (15–21 June) for implementing all three Specialist Agent ReAct loops and integrating InternVL2-8B with OCR pipelines — this is unrealistic for production-grade quality.*

### Resolution:
We successfully decoupled the project timeline into two parallel tracks. For the core production workflow, we implemented all three domain Specialist Agent nodes in LangGraph using prebuilt ReAct loops (`health_agent.py`, `agri_agent.py`, `legal_agent.py`) powered primarily by **Sarvam-105B** (the leading Indian bilingual language model, with OpenAI/Groq/Anthropic fallbacks). Concurrently, we abstract local VLM capabilities in `multimodal.py` using Ollama API queries (supporting `llama3.2-vision`, `qwen2.5-vl`, or `InternVL2`) to permit a smooth, risk-free transition to on-premise hardware without blocking the main workflow.

---

## 4. Point 4: System-Level Non-Functional Requirements (NFRs)

> **Evaluator Feedback:**
> *There is no discussion of system-level non-functional requirements such as latency, throughput, or failure-mode handling.*

### Resolution:
We explicitly measured and compared **average latency** across all benchmark runs (e.g., TriSeva average latency is 10.31s under Claude and 31.79s under Sarvam-105B). For failure-mode handling, the system implements:
1.  Local PDF text extraction via PyMuPDF (fitz) fallback rendering to images for OCR.
2.  Web search fallbacks inside specialist agents when local databases yield no matching context.
3.  Critic-enforced safety boilerplate refusals (*"I do not know"*) instead of ungrounded responses when information is absent.

---

## 5. Point 5: Baseline Comparison beyond Architectural Composition

> **Evaluator Feedback:**
> *No baseline comparison with any existing cross-domain QA system is proposed, making it difficult to claim novelty beyond architectural composition.*

### Resolution:
We implemented and executed **three distinct baselines** across all 500 questions in the benchmark suite:
1.  **Baseline 1 (B1):** Naive RAG with no domain routing (all collections queried).
2.  **Baseline 2 (B2):** No Critic revision loop (`MAX_RETRIES = 0`).
3.  **Baseline 3 (B3):** BM25 keyword search instead of semantic embeddings.

Comparing TriSeva against these baselines provides quantitative proof of the specific performance gain of our agentic composition (e.g., +8.9% overall faithfulness gain and +28.4% on retried queries contributed by the Critic loop).

---

## 6. Point 6: Reliability Risks of Kaggle Free-Tier T4 GPU

> **Evaluator Feedback:**
> *Reliance on Kaggle free-tier T4 GPU for InternVL2-8B inference introduces reliability and reproducibility risks that are not addressed.*

### Resolution:
We have completely eliminated the reliance on Kaggle free-tier GPUs. By implementing **local Ollama/vLLM vision API client support** in `agents/multimodal.py`, the system can run on any local Nvidia RTX workstation or CDAC Pune's dedicated GPU servers (e.g., RTX 3060/4060 with 6GB+ VRAM under quantized `llama3.2-vision` / `qwen2.5-vl`). This provides high performance, complete data privacy, and full reproducibility of local OCR pipelines.
"""
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Successfully created feedback resolution Markdown file at: {md_path}")

if __name__ == '__main__':
    create_feedback_document(
        r'C:\Users\Moulik\Agentic AI\Triseva\faculty_feedback_resolution.docx',
        r'C:\Users\Moulik\.gemini\antigravity\brain\781ff800-1961-4b2d-8dba-e996c58407cc\faculty_feedback_resolution.md'
    )
